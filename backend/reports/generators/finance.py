"""
Finance summary report generator.

One PDF for web and mobile: both ask `POST /api/reports/generate-instant/` for FINANCE_SUMMARY.
The numbers come from the same view code that feeds the Finance Dashboard and Financial Reports
screens (called with the caller's own request), so role rules such as hiding sensitive expenses
and private accounts from PRINCIPAL apply here exactly as they do on screen.
"""

from django.utils import timezone
import io
from datetime import date, datetime
from decimal import Decimal

from django.http import QueryDict

from .base import BaseReportGenerator

MONTH_NAMES = ['Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec']
PRIMARY = '#4F46E5'


def _num(value):
    try:
        return Decimal(str(value if value not in (None, '') else 0))
    except Exception:
        return Decimal('0')


def _fmt(value):
    """1,234 or 1,234.5 (no trailing zeros), like the web's toLocaleString."""
    n = _num(value)
    text = f"{n:,.2f}".rstrip('0').rstrip('.')
    return text or '0'


class _ParamRequest:
    """The caller's request with its query parameters replaced.

    The finance views read their filters from `request.query_params`; everything else (user,
    tenant, role) is delegated to the real request so the same permission filtering applies.
    """

    def __init__(self, request, params):
        self._request = request
        query = QueryDict(mutable=True)
        for key, value in params.items():
            if value not in (None, ''):
                query[key] = str(value)
        self.query_params = query

    def __getattr__(self, name):
        return getattr(self._request, name)


class FinanceSummaryReportGenerator(BaseReportGenerator):
    """Income, expenses, account balances, six-month trend and expense categories for a period."""

    # Set by the report view; this generator reads the live finance views as the caller.
    request = None

    def _call(self, view_cls, action, params):
        view = view_cls()
        proxy = _ParamRequest(self.request, params)
        view.request = proxy
        view.action = action
        view.format_kwarg = None
        view.kwargs = {}
        response = getattr(view, action)(proxy)
        if response.status_code >= 400:
            raise ValueError(f"Finance data unavailable ({action}): {getattr(response, 'data', '')}")
        return response.data

    def get_data(self):
        if self.request is None:
            raise ValueError('FINANCE_SUMMARY needs the requesting user.')
        from finance.views import AccountViewSet, ExpenseViewSet, FeePaymentViewSet, FinanceReportsView

        today = timezone.localdate()
        date_from = self.parameters.get('date_from') or today.replace(day=1).isoformat()
        date_to = self.parameters.get('date_to') or today.isoformat()
        period_label = self.parameters.get('period_label') or 'Custom'
        academic_year = self.parameters.get('academic_year')
        period = {'date_from': date_from, 'date_to': date_to}

        reports_view = FinanceReportsView()
        reports_view.request = _ParamRequest(self.request, period)
        reports_view.format_kwarg = None
        reports_view.kwargs = {}
        school_id = self.school.id
        summary = reports_view._summary(reports_view.request, school_id).data
        trend_view = FinanceReportsView()
        trend_view.request = _ParamRequest(self.request, {'months': 6})
        trend_view.format_kwarg = None
        trend_view.kwargs = {}
        trend = trend_view._monthly_trend(trend_view.request, school_id).data.get('trend', [])

        categories = self._call(ExpenseViewSet, 'category_summary', period)
        accounts = self._call(AccountViewSet, 'balances', {})
        fee = self._call(FeePaymentViewSet, 'monthly_summary', {
            'month': today.month, 'year': today.year, 'fee_type': 'MONTHLY',
            **({'academic_year': academic_year} if academic_year else {}),
        })

        fee_due = _num(fee.get('total_due'))
        fee_collected = _num(fee.get('total_collected'))
        fee_pending = _num(fee.get('total_pending', max(Decimal('0'), fee_due - fee_collected)))
        fee_rate = round(fee_collected / fee_due * 100) if fee_due > 0 else 0

        cat_total = _num(categories.get('total'))
        cat_list = categories.get('categories', [])
        account_list = accounts if isinstance(accounts, list) else []
        grand_total = sum((_num(a.get('net_balance')) for a in account_list), Decimal('0'))

        sections = [{
            'heading': 'Financial Summary',
            'headers': ['Metric', 'Amount'],
            'rows': [
                ['Total Income', _fmt(summary.get('total_income'))],
                ['Total Expenses', _fmt(summary.get('total_expenses'))],
                ['Net Balance', _fmt(summary.get('balance'))],
                ['Fee Collection Rate', f"{fee_rate}%"],
                ['Fee Collected (Current Month)', _fmt(fee_collected)],
                ['Fee Pending (Current Month)', _fmt(fee_pending)],
            ],
            'right': [1],
            'widths': [0.6, 0.4],
        }]
        if account_list:
            sections.append({
                'heading': 'Account Balances',
                'headers': ['Account', 'Type', 'Balance'],
                'rows': [[a.get('name'), a.get('account_type'), _fmt(a.get('net_balance'))] for a in account_list]
                        + [['Total', '', _fmt(grand_total)]],
                'right': [2],
                'widths': [0.45, 0.25, 0.3],
                'bold_last': True,
            })
        if trend:
            sections.append({
                'heading': 'Monthly Trend (Last 6 Months)',
                'headers': ['Month', 'Income', 'Expenses', 'Net Balance'],
                'rows': [
                    [f"{MONTH_NAMES[t['month'] - 1]} {t['year']}", _fmt(t.get('income')), _fmt(t.get('expenses')), _fmt(t.get('balance'))]
                    for t in trend
                ],
                'right': [1, 2, 3],
                'widths': [0.28, 0.24, 0.24, 0.24],
            })
        if cat_list:
            sections.append({
                'heading': 'Expense Breakdown',
                'headers': ['Category', 'Amount', '% of Total', 'Entries'],
                'rows': [
                    [c.get('category_display'), _fmt(c.get('total_amount')),
                     f"{round(_num(c.get('total_amount')) / cat_total * 100)}%" if cat_total > 0 else '0%', c.get('count')]
                    for c in cat_list
                ] + [['Total', _fmt(cat_total), '100%', '']],
                'right': [1, 2],
                'widths': [0.35, 0.25, 0.22, 0.18],
                'bold_last': True,
            })

        return {
            'title': 'Finance Report',
            'subtitle': f"Period: {period_label} ({date_from} to {date_to})",
            'sections': sections,
        }

    def render_pdf(self, data):
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4, leftMargin=14 * mm, rightMargin=14 * mm,
                                topMargin=14 * mm, bottomMargin=16 * mm)
        width = A4[0] - 28 * mm
        styles = getSampleStyleSheet()
        title = ParagraphStyle('FinTitle', parent=styles['Heading1'], fontSize=20, textColor=colors.HexColor(PRIMARY), spaceAfter=4)
        heading = ParagraphStyle('FinHeading', parent=styles['Heading3'], fontSize=13, spaceBefore=14, spaceAfter=6)
        muted = ParagraphStyle('FinMuted', parent=styles['Normal'], fontSize=9, textColor=colors.HexColor('#646464'))

        elements = [
            Paragraph(data['title'], title),
            Paragraph(self.school.name, styles['Normal']),
            Paragraph(data['subtitle'], muted),
            Paragraph(f"Generated: {timezone.localtime().strftime('%B %d, %Y %I:%M %p')}", muted),
            Spacer(1, 8),
        ]
        for section in data['sections']:
            elements.append(Paragraph(section['heading'], heading))
            rows = [section['headers']] + [['' if c is None else str(c) for c in r] for r in section['rows']]
            table = Table(rows, colWidths=[width * w for w in section['widths']], repeatRows=1)
            style = [
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor(PRIMARY)),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, -1), 9),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F3F4F6')]),
                ('TOPPADDING', (0, 0), (-1, -1), 4),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ]
            style += [('ALIGN', (c, 0), (c, -1), 'RIGHT') for c in section.get('right', [])]
            if section.get('bold_last'):
                style.append(('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'))
            table.setStyle(TableStyle(style))
            elements.append(table)

        def footer(canvas, document):
            canvas.saveState()
            canvas.setFont('Helvetica', 8)
            canvas.setFillColor(colors.HexColor('#969696'))
            canvas.drawString(14 * mm, 9 * mm, 'Education AI Finance Report')
            canvas.drawRightString(A4[0] - 14 * mm, 9 * mm, f"Page {document.page}")
            canvas.restoreState()

        doc.build(elements, onFirstPage=footer, onLaterPages=footer)
        return buffer.getvalue()

    def render_excel(self, data):
        from openpyxl import Workbook
        from openpyxl.styles import Font

        wb = Workbook()
        ws = wb.active
        ws.title = 'Finance Report'
        ws['A1'] = self.school.name
        ws['A1'].font = Font(size=14, bold=True)
        ws['A2'] = data['title']
        ws['A3'] = data['subtitle']
        row = 5
        for section in data['sections']:
            ws.cell(row=row, column=1, value=section['heading']).font = Font(bold=True)
            row += 1
            for col, header in enumerate(section['headers'], 1):
                ws.cell(row=row, column=col, value=header).font = Font(bold=True)
            for r in section['rows']:
                row += 1
                for col, value in enumerate(r, 1):
                    ws.cell(row=row, column=col, value=value)
            row += 2
        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()
